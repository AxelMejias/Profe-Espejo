from io import BytesIO
from typing import Annotated, Optional
import math
from fastapi import APIRouter, Depends, Query, Path, Body, UploadFile, File, status, HTTPException  # noqa: F401
from fastapi.responses import StreamingResponse
import openpyxl

from app.modules.productos.schemas import (
    ProductoCreate, ProductoUpdate, ProductoRead, PaginatedProductos,
    ImagenProductoUpdate, AsociarIngredienteRequest, ProductoIngredienteRead,
    InsumoEnProductoRead,
)
from app.modules.productos import service
from app.core.dependencies import require_role
from app.core.unit_of_work import UnitOfWork
from app.core.websocket import emit_catalogo_evento


router = APIRouter(prefix="/api/v1/productos", tags=["Productos"])

# Catálogo de lectura: PÚBLICO según doc §5.2 (GET /productos y /productos/{id}
# son navegables sin autenticación).
_DISPONIBILIDAD = Depends(require_role(["ADMIN", "STOCK"]))   # toggle disponibilidad
_ADMIN          = Depends(require_role(["ADMIN"]))             # crear, editar, borrar, importar, reactivar


@router.get("/exportar", summary="Exportar productos activos a Excel")
def exportar_productos(_=_ADMIN):
    rows = []
    with UnitOfWork() as uow:
        items, _ = uow.productos.get_all(solo_disponibles=False, page=1, size=10000)
        for prod in items:
            cats = ", ".join(c.nombre for c in prod.categorias)
            links = uow.productos.get_ingrediente_links(prod.id)
            insumos_parts = []
            for link in links:
                ing = uow.ingredientes.get_by_id(link.ingrediente_id)
                if ing:
                    insumos_parts.append(f"{ing.nombre}:{link.cantidad}")
            rows.append([
                prod.id, prod.nombre, prod.descripcion or "",
                float(prod.precio_base), float(prod.margen_ganancia) * 100,
                "Sí" if prod.disponible else "No",
                cats, ", ".join(insumos_parts),
            ])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Productos"
    ws.append(["ID", "Nombre", "Descripción", "Precio", "Margen (%)", "Disponible", "Categorías", "Insumos"])
    for row in rows:
        ws.append(row)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=productos.xlsx"},
    )


@router.get("/plantilla", summary="Descargar plantilla Excel para importar productos")
def descargar_plantilla(_=_ADMIN):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Productos"
    ws.append([
        "Nombre", "Descripcion", "Margen Ganancia (ej: 0.30)",
        "Disponible", "Categorias (nombres separados por coma)",
        "Insumos (nombre:cantidad separados por coma)",
    ])
    ws.append([
        "Hamburguesa Clásica", "", "0.30", "TRUE",
        "Comidas rápidas", "Carne:0.2,Pan:1,Lechuga:0.05",
    ])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla_productos.xlsx"},
    )


@router.post("/importar", summary="Importar productos desde Excel")
async def importar_productos(archivo: UploadFile = File(...), _=_ADMIN):
    contents = archivo.file.read()
    wb = openpyxl.load_workbook(BytesIO(contents), data_only=True)
    ws = wb.active

    creados = 0
    omitidos = 0
    errores = []

    # El router solo orquesta: lee el archivo, abre una transacción por fila y
    # delega la lógica de negocio (validación, parseo, precio, alta) al service.
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or all(v is None for v in row):
            continue
        try:
            with UnitOfWork() as uow:
                resultado = service.importar_fila(uow, row)
            if resultado == "creado":
                creados += 1
            else:
                omitidos += 1
        except Exception as e:
            nombre_str = str(row[0]) if row and row[0] is not None else ""
            errores.append({"fila": idx, "nombre": nombre_str, "motivo": str(e)})

    if creados:
        await emit_catalogo_evento("producto_actualizado")
    return {"creados": creados, "omitidos": omitidos, "errores": errores}



@router.get("/", response_model=PaginatedProductos, summary="Listar productos")
def listar_productos(
    nombre:           Annotated[Optional[str],   Query(max_length=100)] = None,
    precio_min:       Annotated[Optional[float], Query(ge=0)]           = None,
    precio_max:       Annotated[Optional[float], Query(ge=0)]           = None,
    categoria_id:     Annotated[Optional[int],   Query(ge=1)]           = None,
    solo_disponibles: Annotated[bool, Query()]                          = True,
    solo_destacados:  Annotated[bool, Query()]                          = False,
    con_stock:        Annotated[Optional[bool], Query()]                = None,
    page:             Annotated[int, Query(ge=1)]                       = 1,
    size:             Annotated[int, Query(ge=1, le=100)]               = 20,
):
    with UnitOfWork() as uow:
        return service.get_all(
            uow, nombre=nombre, precio_min=precio_min, precio_max=precio_max,
            categoria_id=categoria_id, solo_disponibles=solo_disponibles,
            solo_destacados=solo_destacados, con_stock=con_stock, page=page, size=size,
        )


@router.get("/inactivos", response_model=PaginatedProductos, summary="Listar productos inactivos")
def listar_productos_inactivos(
    page: Annotated[int, Query(ge=1)] = 1,
    size: Annotated[int, Query(ge=1, le=100)] = 20,
    _=_ADMIN,
):
    with UnitOfWork() as uow:
        items, total = uow.productos.get_all_inactivos(page=page, size=size)
        return PaginatedProductos(
            items=[service._build_response(uow, p) for p in items],
            total=total, page=page, size=size,
            pages=math.ceil(total / size) if total else 0,
        )


@router.get("/{producto_id}", response_model=ProductoRead, summary="Obtener producto por ID")
def obtener_producto(producto_id: Annotated[int, Path(ge=1)]):
    with UnitOfWork() as uow:
        return service.get_by_id(uow, producto_id)


@router.post("/", response_model=ProductoRead, status_code=status.HTTP_201_CREATED, summary="Crear producto")
async def crear_producto(data: ProductoCreate, _=_ADMIN):
    with UnitOfWork() as uow:
        result = service.create(uow, data)
    await emit_catalogo_evento("producto_creado", producto_id=result.id)
    return result


@router.put("/{producto_id}", response_model=ProductoRead, summary="Actualizar producto")
async def actualizar_producto(
    producto_id: Annotated[int, Path(ge=1)],
    data: ProductoUpdate,
    _=_ADMIN,
):
    with UnitOfWork() as uow:
        result = service.update(uow, producto_id, data)
    await emit_catalogo_evento("producto_actualizado", producto_id=result.id)
    return result


@router.patch("/{producto_id}/reactivar", response_model=ProductoRead, summary="Reactivar producto inactivo")
async def reactivar_producto(producto_id: Annotated[int, Path(ge=1)], _=_ADMIN):
    with UnitOfWork() as uow:
        result = service.reactivar(uow, producto_id)
    await emit_catalogo_evento("producto_actualizado", producto_id=result.id)
    return result


@router.patch("/{producto_id}/disponibilidad", response_model=ProductoRead, summary="Toggle disponibilidad")
async def toggle_disponibilidad(
    producto_id: Annotated[int, Path(ge=1)],
    disponible: bool = Body(..., embed=True),
    _=_DISPONIBILIDAD,
):
    with UnitOfWork() as uow:
        result = service.toggle_disponibilidad(uow, producto_id, disponible)
    await emit_catalogo_evento("producto_actualizado", producto_id=result.id)
    return result


@router.patch("/{producto_id}/imagenes", response_model=ProductoRead, summary="Actualizar lista de imágenes del producto")
async def actualizar_imagenes(
    producto_id: Annotated[int, Path(ge=1)],
    data: ImagenProductoUpdate,
    _=_ADMIN,
):
    with UnitOfWork() as uow:
        result = service.set_imagenes(uow, producto_id, data.imagenes_url)
    await emit_catalogo_evento("producto_actualizado", producto_id=result.id)
    return result


@router.get("/{producto_id}/ingredientes", response_model=list[InsumoEnProductoRead], summary="Listar insumos del producto")
def listar_ingredientes_producto(producto_id: Annotated[int, Path(ge=1)]):
    with UnitOfWork() as uow:
        return service.listar_ingredientes(uow, producto_id)


@router.post("/{producto_id}/ingredientes", response_model=ProductoIngredienteRead,
             status_code=status.HTTP_201_CREATED, summary="Asociar insumo al producto")
async def asociar_ingrediente(
    producto_id: Annotated[int, Path(ge=1)],
    data: AsociarIngredienteRequest,
    _=_ADMIN,
):
    with UnitOfWork() as uow:
        result = service.agregar_ingrediente(uow, producto_id, data)
    await emit_catalogo_evento("producto_actualizado", producto_id=producto_id)
    return result


@router.patch("/{producto_id}/destacar", response_model=ProductoRead, summary="Toggle destacado en Home")
async def toggle_destacado(
    producto_id: Annotated[int, Path(ge=1)],
    destacado: bool = Body(..., embed=True),
    _=_ADMIN,
):
    with UnitOfWork() as uow:
        result = service.toggle_destacado(uow, producto_id, destacado)
    await emit_catalogo_evento("producto_actualizado", producto_id=result.id)
    return result


@router.delete("/{producto_id}", status_code=status.HTTP_204_NO_CONTENT, summary="Baja lógica de producto")
async def eliminar_producto(producto_id: Annotated[int, Path(ge=1)], _=_ADMIN):
    with UnitOfWork() as uow:
        service.delete(uow, producto_id)
    await emit_catalogo_evento("producto_eliminado", producto_id=producto_id)
