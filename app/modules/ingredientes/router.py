from io import BytesIO
from typing import Annotated, Optional
from fastapi import APIRouter, Depends, Query, Path, UploadFile, File, status
from fastapi.responses import StreamingResponse
import openpyxl

from app.modules.ingredientes.schemas import (
    IngredienteCreate, IngredienteUpdate, IngredienteResponse, PaginatedIngredientes,
)
from app.modules.ingredientes import service
from app.core.dependencies import require_role
from app.core.unit_of_work import UnitOfWork

router = APIRouter(prefix="/api/v1/ingredientes", tags=["Ingredientes"])

_LEER       = Depends(require_role(["ADMIN", "STOCK", "PEDIDOS", "CLIENT"]))
_STOCK_EDIT = Depends(require_role(["ADMIN", "STOCK"]))   # leer + actualizar stock vía PUT
_ADMIN      = Depends(require_role(["ADMIN"]))             # crear, borrar, importar, reactivar

_BOOL_MAP = {"TRUE", "VERDADERO", "SI", "SÍ", "S", "1"}


@router.get("/exportar", summary="Exportar ingredientes activos a Excel")
def exportar_ingredientes(_=_STOCK_EDIT):
    with UnitOfWork() as uow:
        ingredientes = service.get_all_activos_for_export(uow)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ingredientes"
    ws.append(["ID", "Nombre", "Descripción", "Unidad de medida", "Es alérgeno", "Es terminado", "Creado en"])
    for ing in ingredientes:
        ws.append([
            ing.id, ing.nombre, ing.descripcion or "",
            ing.unidad_medida, "Sí" if ing.es_alergeno else "No",
            "Sí" if ing.es_producto_terminado else "No",
            ing.created_at.strftime("%Y-%m-%d %H:%M"),
        ])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=ingredientes.xlsx"},
    )


@router.get("/plantilla", summary="Descargar plantilla Excel para importar ingredientes")
def descargar_plantilla(_=_ADMIN):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ingredientes"
    ws.append([
        "Nombre", "Descripcion", "Unidad Medida", "Costo Unitario",
        "Stock Actual", "Stock Minimo", "Es Alergeno", "Es Producto Terminado",
    ])
    ws.append(["Tomate", "Tomate fresco", "kg", 150.00, 10.0, 2.0, "FALSE", "FALSE"])
    ws.append(["Coca Cola", "Bebida 500ml", "unidad", 800.00, 24.0, 6.0, "FALSE", "TRUE"])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla_ingredientes.xlsx"},
    )


@router.post("/importar", summary="Importar ingredientes desde Excel")
def importar_ingredientes(archivo: UploadFile = File(...), _=_ADMIN):
    from app.modules.ingredientes.model import Ingrediente

    contents = archivo.file.read()
    wb = openpyxl.load_workbook(BytesIO(contents), data_only=True)
    ws = wb.active

    creados = 0
    omitidos = 0
    errores = []

    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or all(v is None for v in row):
            continue
        try:
            nombre = str(row[0]).strip() if row[0] is not None else ""
            descripcion = str(row[1]).strip() if len(row) > 1 and row[1] is not None else None
            unidad_medida = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
            costo_raw = row[3] if len(row) > 3 else None
            stock_raw = row[4] if len(row) > 4 else 0
            minimo_raw = row[5] if len(row) > 5 else 0
            alergeno_raw = str(row[6]).upper().strip() if len(row) > 6 and row[6] is not None else "FALSE"
            terminado_raw = str(row[7]).upper().strip() if len(row) > 7 and row[7] is not None else "FALSE"

            if not nombre:
                errores.append({"fila": idx, "nombre": "", "motivo": "Nombre requerido"})
                continue
            if not unidad_medida:
                errores.append({"fila": idx, "nombre": nombre, "motivo": "Unidad de medida requerida"})
                continue
            if costo_raw is None:
                errores.append({"fila": idx, "nombre": nombre, "motivo": "Costo unitario requerido"})
                continue

            costo = float(costo_raw)
            if costo < 0:
                errores.append({"fila": idx, "nombre": nombre, "motivo": "Costo unitario debe ser >= 0"})
                continue

            with UnitOfWork() as uow:
                if uow.ingredientes.get_by_nombre(nombre):
                    omitidos += 1
                    continue
                ing = Ingrediente(
                    nombre=nombre,
                    descripcion=descripcion or None,
                    unidad_medida=unidad_medida,
                    costo_unitario=costo,
                    stock_cantidad=float(stock_raw) if stock_raw is not None else 0,
                    stock_minimo=float(minimo_raw) if minimo_raw is not None else 0,
                    es_alergeno=alergeno_raw in _BOOL_MAP,
                    es_producto_terminado=terminado_raw in _BOOL_MAP,
                )
                uow.ingredientes.add(ing)
                creados += 1
        except Exception as e:
            nombre_str = str(row[0]) if row and row[0] is not None else ""
            errores.append({"fila": idx, "nombre": nombre_str, "motivo": str(e)})

    return {"creados": creados, "omitidos": omitidos, "errores": errores}


@router.get("/", response_model=PaginatedIngredientes, summary="Listar ingredientes")
def listar_ingredientes(
    nombre:                Annotated[Optional[str],  Query(max_length=100)] = None,
    unidad_medida:         Annotated[Optional[str],  Query(max_length=50)]  = None,
    es_alergeno:           Annotated[Optional[bool], Query()]               = None,
    es_producto_terminado: Annotated[Optional[bool], Query()]               = None,
    page:                  Annotated[int, Query(ge=1)]                      = 1,
    size:                  Annotated[int, Query(ge=1, le=100)]              = 20,
    _=_LEER,
):
    with UnitOfWork() as uow:
        return service.get_all(
            uow, nombre=nombre, es_alergeno=es_alergeno,
            es_producto_terminado=es_producto_terminado,
            unidad_medida=unidad_medida, page=page, size=size,
        )


@router.get(
    "/inactivos",
    response_model=PaginatedIngredientes,
    summary="Listar ingredientes dados de baja (soft delete)",
)
def listar_ingredientes_inactivos(
    page: Annotated[int, Query(ge=1)] = 1,
    size: Annotated[int, Query(ge=1, le=100)] = 20,
    _=_ADMIN,
):
    with UnitOfWork() as uow:
        items, total = uow.ingredientes.get_all_inactivos(page=page, size=size)
        import math
        return PaginatedIngredientes(
            items=[IngredienteResponse.model_validate(i) for i in items],
            total=total, page=page, size=size,
            pages=math.ceil(total / size) if total else 0,
        )


@router.get("/{ingrediente_id}", response_model=IngredienteResponse, summary="Obtener ingrediente por ID")
def obtener_ingrediente(ingrediente_id: Annotated[int, Path(ge=1)], _=_LEER):
    with UnitOfWork() as uow:
        return service.get_by_id(uow, ingrediente_id)


@router.post("/", response_model=IngredienteResponse, status_code=status.HTTP_201_CREATED, summary="Crear ingrediente")
def crear_ingrediente(data: IngredienteCreate, _=_ADMIN):
    with UnitOfWork() as uow:
        return service.create(uow, data)


@router.put("/{ingrediente_id}", response_model=IngredienteResponse, summary="Actualizar ingrediente")
def actualizar_ingrediente(
    ingrediente_id: Annotated[int, Path(ge=1)],
    data: IngredienteUpdate,
    _=_STOCK_EDIT,
):
    with UnitOfWork() as uow:
        return service.update(uow, ingrediente_id, data)


@router.patch("/{ingrediente_id}/reactivar", response_model=IngredienteResponse, summary="Reactivar ingrediente dado de baja")
def reactivar_ingrediente(ingrediente_id: Annotated[int, Path(ge=1)], _=_ADMIN):
    with UnitOfWork() as uow:
        return service.reactivar(uow, ingrediente_id)


@router.delete("/{ingrediente_id}", status_code=status.HTTP_204_NO_CONTENT, summary="Baja lógica de ingrediente")
def eliminar_ingrediente(ingrediente_id: Annotated[int, Path(ge=1)], _=_ADMIN):
    with UnitOfWork() as uow:
        service.delete(uow, ingrediente_id)
